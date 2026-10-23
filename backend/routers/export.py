"""Excel export router — tüm veriyi multi-sheet XLSX olarak dışa aktarır."""
import io
import os
from html import escape
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, A3, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from database import db
from dependencies import get_current_user, write_audit

router = APIRouter(prefix="/export", tags=["export"])

REPORT_NAMES = {
    "by-ship-detail": "Gemi Bazlı Borç", "aging": "Borç Yaşlandırma",
    "monthly-projection": "Aylık Projeksiyon", "top-vendors": "Top Tedarikçiler",
    "by-currency": "Döviz Pozisyonu", "cash-report": "Kasa Raporu",
    "office-payments": "Günlük Ofis Ödemeleri", "all-debts": "Tüm Borçlar",
}


def _header_style(ws, row=1):
    """Apple/Notion stili minimal header."""
    for cell in ws[row]:
        cell.font = Font(bold=True, color="86868B", size=10, name="Helvetica")
        cell.fill = PatternFill("solid", fgColor="FAFAFA")
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = Border(bottom=Side(style="thin", color="E5E5EA"))


def _auto_width(ws, max_w=50):
    """Sütun genişliklerini içeriğe göre otomatik ayarla."""
    for col_cells in ws.columns:
        col_letter = col_cells[0].column_letter
        max_len = 10
        for c in col_cells[:200]:
            try:
                if c.value is not None:
                    max_len = max(max_len, min(len(str(c.value)) + 2, max_w))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max_len


def _write_sheet(wb, name, rows, columns):
    """rows: list[dict], columns: list[(field, label)]"""
    ws = wb.create_sheet(name[:31])  # Excel sheet name max 31 char
    headers = [label for _, label in columns]
    ws.append(headers)
    _header_style(ws, 1)
    for r in rows:
        ws.append([r.get(field) for field, _ in columns])
    _auto_width(ws)
    return ws


async def _report_tables(report_id, start=None, end=None, months=30, user=None):
    """Return [(title, headers, rows)] for screen reports and both export formats."""
    from routers import reports
    if report_id == "cash-report":
        data = await reports.cash_report_data(start, end)
        movements = [[p.get("id"), p.get("type"), p.get("date"), p.get("description") or p.get("vendor"),
            p.get("payment_method"), p.get("paying_company"), p.get("currency"), p.get("fx_rate"),
            p.get("amount"), p.get("usd_amount"), p.get("ship")] for p in data["movements"]]
        balances = [[b.get("name"), b.get("type"), b.get("balance_currency"), b.get("opening_balance"),
            b.get("in"), b.get("out"), b.get("current_balance")] for b in data["balances"]]
        return [("Hareketler", ["No", "İşlem Türü", "Tarih", "Açıklama", "Ödeme Şekli", "Şirket", "Döviz", "Kur", "Tutar", "USD", "Masraf Merkezi"], movements),
                ("Hesap Bakiyeleri", ["Hesap", "Tip", "Para Birimi", "Açılış", "Giren USD", "Çıkan USD", "Mevcut Bakiye"], balances)]
    if report_id == "office-payments":
        data = await reports.office_payments_data(start, end)
        rows = []
        for week in data["weeks"]:
            for p in week["rows"]:
                rows.append([week["week"], week["start"], week["end"], p["date"], p["description"], p["vendor"],
                    p["company_ship"], p["amount"], p["currency"], p["usd_amount"], p["budget_category"]])
        return [("Ofis Ödemeleri", ["Hafta", "Başlangıç", "Bitiş", "Tarih", "Açıklama", "Firma", "Şirket/Gemi", "Tutar", "Döviz", "USD", "Bütçe Sınıfı"], rows)]
    if report_id == "all-debts":
        data = await reports.all_debts_data(months)
        headers = ["Grup", "Açıklama/Firma", "Gemi", "Borç Türü", "Vade", "Vadesi Geçmiş", "Toplam USD", *data["months"]]
        rows = [[p["group"], p["description"], p["ship"], p["debt_type"], p["due_date"], p["overdue"], p["total"],
                 *[p["monthly"].get(month, 0) for month in data["months"]]] for p in data["rows"]]
        return [("Tüm Borçlar", headers, rows)]

    if report_id == "by-ship-detail":
        data = await reports.report_by_ship(user)
        return [(REPORT_NAMES[report_id], ["Gemi", "Toplam Borç", "Ödenen", "Açık", "Kayıt"], [[r.get("_id"), r.get("total_debt"), r.get("paid"), r.get("open"), r.get("count")] for r in data])]
    if report_id == "aging":
        data = await reports.report_aging(user)
        return [(REPORT_NAMES[report_id], ["Vade Grubu", "Toplam USD", "Kayıt"], [[r.get("bucket"), r.get("total"), r.get("count")] for r in data])]
    if report_id == "monthly-projection":
        data = await reports.report_monthly(months, user)
        return [(REPORT_NAMES[report_id], ["Ay", "Toplam USD", "Kayıt"], [[r.get("month"), r.get("total"), r.get("count")] for r in data])]
    if report_id == "top-vendors":
        data = await reports.report_top_vendors(20, user)
        return [(REPORT_NAMES[report_id], ["Tedarikçi", "Toplam USD", "İşlem"], [[r.get("vendor"), r.get("total"), r.get("count")] for r in data])]
    if report_id == "by-currency":
        data = await reports.report_currency_position(user)
        return [(REPORT_NAMES[report_id], ["Para Birimi", "Borç", "Alacak", "Net"], [[r.get("currency"), r.get("borç"), r.get("alacak"), (r.get("alacak") or 0) - (r.get("borç") or 0)] for r in data])]
    raise HTTPException(404, "Rapor bulunamadı.")


def _report_xlsx(title, tables):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_title, headers, rows in tables:
        ws = wb.create_sheet(sheet_title[:31])
        ws.append(headers)
        _header_style(ws)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        _auto_width(ws, 35)
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _pdf_font():
    name = "Helvetica"
    paths = [
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        os.path.join(os.path.dirname(__import__("reportlab").__file__), "fonts", "Vera.ttf"),
    ]
    path = next((candidate for candidate in paths if os.path.exists(candidate)), None)
    if path:
        try:
            pdfmetrics.registerFont(TTFont("MartiUnicode", path))
            name = "MartiUnicode"
        except Exception:
            pass
    return name


def _report_pdf(title, tables):
    stream = io.BytesIO()
    page = landscape(A3) if any(len(headers) > 10 for _, headers, _ in tables) else landscape(A4)
    doc = SimpleDocTemplate(stream, pagesize=page, leftMargin=22, rightMargin=22, topMargin=24, bottomMargin=24)
    font = _pdf_font()
    styles = getSampleStyleSheet()
    styles["Title"].fontName = font
    styles["Heading2"].fontName = font
    styles["Normal"].fontName = font
    styles["Normal"].fontSize = 7
    styles["Normal"].leading = 9
    story = [Paragraph(title, styles["Title"]), Paragraph(f"Oluşturulma: {datetime.now().strftime('%d.%m.%Y %H:%M')}", styles["Normal"]), Spacer(1, 12)]
    first_output = True
    for sheet_title, headers, rows in tables:
        segments = [(sheet_title, headers, rows)]
        if len(headers) > 15:
            fixed, chunk = 7, 8
            segments = []
            for offset in range(fixed, len(headers), chunk):
                indexes = [*range(fixed), *range(offset, min(offset + chunk, len(headers)))]
                segments.append((f"{sheet_title} · {headers[offset]} – {headers[indexes[-1]]}",
                    [headers[i] for i in indexes], [[row[i] for i in indexes] for row in rows]))
        for segment_title, segment_headers, segment_rows in segments:
            if not first_output:
                story.append(PageBreak())
            first_output = False
            story.append(Paragraph(segment_title, styles["Heading2"]))
            story.append(Spacer(1, 6))
            cell = lambda value: Paragraph(escape(str(value if value is not None else "—")), styles["Normal"])
            data = [[cell(h) for h in segment_headers], *[[cell(v) for v in row] for row in segment_rows]]
            available = page[0] - 44
            widths = [available / max(1, len(segment_headers))] * len(segment_headers)
            table = Table(data, repeatRows=1, colWidths=widths)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111111")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("FONTSIZE", (0, 0), (-1, -1), 6 if len(segment_headers) > 10 else 7),
                ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#E5E5EA")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFAFA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(table)
    doc.build(story)
    stream.seek(0)
    return stream


@router.get("/excel")
async def export_excel(user: dict = Depends(get_current_user)):
    """Tüm veriyi Excel olarak indirir — 1 dosyada tüm sheet'ler."""
    wb = Workbook()
    wb.remove(wb.active)  # default sheet'i sil

    # 1. Borçlar
    payables = await db.payables.find({"kind": "PAYABLE"}, {"_id": 0}).sort("due_date", -1).to_list(10000)
    _write_sheet(wb, "Borçlar", payables, [
        ("due_date", "Vade"), ("vendor", "Tedarikçi"), ("ship", "Gemi/Birim"),
        ("expense_type", "Masraf Türü"), ("expense_code", "Masraf Kodu"),
        ("budget_category", "Bütçe Sınıfı"),
        ("description", "Açıklama"), ("original_amount", "Tutar"),
        ("currency", "Döviz"), ("usd_amount", "USD Karşılığı"),
        ("status", "Durum"), ("armator", "Armatör"),
        ("country", "Ülke"), ("year", "Yıl"), ("month", "Ay"),
        ("created_at", "Eklenme Tarihi"), ("created_by", "Ekleyen"),
    ])

    # 2. Alacaklar
    receivables = await db.payables.find({"kind": "RECEIVABLE"}, {"_id": 0}).sort("due_date", -1).to_list(10000)
    _write_sheet(wb, "Alacaklar", receivables, [
        ("due_date", "Vade"), ("vendor", "Müşteri"), ("ship", "Gemi/Birim"),
        ("description", "Açıklama"), ("original_amount", "Tutar"),
        ("currency", "Döviz"), ("usd_amount", "USD Karşılığı"),
        ("status", "Durum"), ("created_at", "Eklenme"),
    ])

    # 3. Ödemeler/Tahsilatlar
    payments = await db.payments.find({}, {"_id": 0}).sort("date", -1).to_list(10000)
    _write_sheet(wb, "Ödemeler", payments, [
        ("date", "Tarih"), ("type", "Tip"), ("vendor", "Firma"),
        ("description", "Açıklama"), ("paying_company", "Ödeyen Şirket"),
        ("payment_method", "Banka/Kasa"), ("ship", "Gemi/Birim"),
        ("amount", "Tutar"), ("currency", "Döviz"), ("fx_rate", "Kur"),
        ("usd_amount", "USD Karşılığı"), ("approved", "Onaylı"), ("posting_status", "Kayıt Durumu"), ("voided", "İptal"), ("payable_id", "Bağlı Belge"), ("fx_difference_usd", "Kur Farkı (USD)"),
        ("created_at", "Eklenme"), ("created_by", "Ekleyen"),
    ])

    # 4. Cari Hesaplar Özet
    from routers.accounts import _vendor_ledgers
    ca_rows = sorted(await _vendor_ledgers(db), key=lambda x: x['balance'], reverse=True)
    _write_sheet(wb, "Cari Hesaplar", ca_rows, [
        ("name", "Cari Adı"), ("debt", "Toplam Borç (USD)"), ("paid", "Ödenen Mahsup (USD)"),
        ("receivable", "Toplam Alacak (USD)"), ("received", "Tahsil Mahsup (USD)"),
        ("cash_paid", "Nakit Ödeme (USD)"), ("cash_received", "Nakit Tahsil (USD)"),
        ("balance", "Net Borç / Alacak (USD)"), ("count", "Belge Sayısı")])

    # 5. Master Data — Tedarikçiler
    vendors = await db.vendors.find({}, {"_id": 0}).to_list(5000)
    _write_sheet(wb, "Tedarikçiler", vendors, [
        ("name", "Ad"), ("country", "Ülke"), ("tax_no", "Vergi No"),
        ("iban", "IBAN"), ("contact", "İletişim"), ("phone", "Telefon"),
        ("email", "Email"), ("notes", "Notlar"),
    ])

    # 6. Gemiler
    ships = await db.ships.find({}, {"_id": 0}).to_list(500)
    _write_sheet(wb, "Gemiler", ships, [
        ("name", "Gemi/Birim"), ("manager", "Manager"), ("armator", "Armatör"),
        ("imo", "IMO"), ("flag", "Bayrak"), ("notes", "Notlar"),
    ])

    # 7. Bankalar
    banks = await db.banks.find({}, {"_id": 0}).to_list(500)
    _write_sheet(wb, "Bankalar", banks, [
        ("name", "Banka/Kasa Adı"), ("type", "Tip"), ("currency", "Para Birimi"),
        ("iban", "IBAN"), ("balance", "Bakiye"), ("company", "Şirket"),
    ])

    # 8. Şirketler
    companies = await db.companies.find({}, {"_id": 0}).to_list(500)
    _write_sheet(wb, "Şirketler", companies, [
        ("name", "Şirket Adı"), ("tax_no", "Vergi No"), ("notes", "Notlar"),
    ])

    # 9. Masraf Türleri
    exp_types = await db.expense_types.find({}, {"_id": 0}).to_list(500)
    _write_sheet(wb, "Masraf Türleri", exp_types, [
        ("code", "Kod"), ("name", "Ad"), ("notes", "Notlar"),
    ])

    # 10. Muhasebe Hesap Planı
    acc_codes = await db.accounting_codes.find({}, {"_id": 0}).to_list(500)
    _write_sheet(wb, "Hesap Planı", acc_codes, [
        ("code", "Kod"), ("name", "Hesap Adı"),
    ])

    # 11. Döviz Kurları
    fx = await db.currencies.find({}, {"_id": 0}).to_list(50)
    _write_sheet(wb, "Döviz Kurları", fx, [
        ("code", "Kod"), ("name", "Ad"), ("rate_to_tl", "TL Karşılığı"),
        ("forex_selling", "Döviz Satış"), ("banknote_buying", "Efektif Alış"),
        ("banknote_selling", "Efektif Satış"), ("last_updated", "Son Güncelleme"),
    ])

    # 12. Armatörler
    armators = await db.armators.find({}, {"_id": 0}).to_list(500)
    _write_sheet(wb, "Armatörler", armators, [("name", "Ad"), ("notes", "Notlar")])

    # 13. Kişi & Alt Şirketler
    people = await db.people.find({}, {"_id": 0}).to_list(500)
    _write_sheet(wb, "Kişi & Şirketler", people, [("name", "Ad"), ("notes", "Notlar")])

    # 14. Ülkeler
    countries = await db.countries.find({}, {"_id": 0}).to_list(500)
    _write_sheet(wb, "Ülkeler", countries, [("name", "Ad"), ("code", "Kod")])

    # Stream as response
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"marti-finans-veri-{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"
    await write_audit(user, "export", "excel", meta={"filename": filename})
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/reports/{report_id}")
async def export_report(report_id: str, format: str = "xlsx", start: str | None = None,
                        end: str | None = None, months: int = 30,
                        user: dict = Depends(get_current_user)):
    """Download any report as a styled XLSX or printable PDF."""
    if format not in {"xlsx", "pdf"}:
        raise HTTPException(400, "Format xlsx veya pdf olmalıdır.")
    tables = await _report_tables(report_id, start, end, months, user)
    title = REPORT_NAMES.get(report_id, report_id)
    stream = _report_xlsx(title, tables) if format == "xlsx" else _report_pdf(title, tables)
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    filename = f"marti-{report_id}-{stamp}.{format}"
    await write_audit(user, "export", f"report-{format}", meta={"report": report_id, "filename": filename})
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if format == "xlsx" else "application/pdf"
    return StreamingResponse(stream, media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}"'})
